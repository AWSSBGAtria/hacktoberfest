#!/usr/bin/env python3
"""Write a one-row-per-person version of the final shortlist sheet.

hacktoberfest_teams_final_formatted.xlsx next to the source file:
- one person per row (no semicolon-joined cells)
- team name merged across its members' rows; blank for solo entries
- canonical college names; blank when the sheet has no college
- solos last, teams in the same order the site uses (case-insensitive)
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_shortlisted import NOT_SPECIFIED, XLSX, load_roster

OUT = XLSX.replace('.xlsx', '_formatted.xlsx')

THIN = Side(style='thin', color='C9C9C0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main():
    teams, solos, notes = load_roster()
    for note in notes:
        print('note:', note)
    by_name = lambda g: g['team'].casefold()
    teams.sort(key=by_name)
    solos.sort(key=by_name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Attendees'

    ws.append(['Team Name', 'Member Name', 'Email', 'College / Organization'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='F7F7F2')
        cell.fill = PatternFill('solid', fgColor='10201D')
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = BORDER
    ws.row_dimensions[1].height = 20

    row = 2
    for group in teams + solos:
        start = row
        for member in group['members']:
            if row == start:
                ws.cell(row=row, column=1, value=None if group['solo'] else group['team'])
            ws.cell(row=row, column=2, value=member['name'])
            ws.cell(row=row, column=3, value=member['email'])
            college = member['college']
            ws.cell(row=row, column=4,
                    value=None if college == NOT_SPECIFIED else college)
            for col in range(1, 5):
                cell = ws.cell(row=row, column=col)
                cell.font = Font(size=11)
                cell.border = BORDER
                cell.alignment = Alignment(vertical='center')
            row += 1
        if not group['solo'] and row - start > 1:
            ws.merge_cells(start_row=start, start_column=1,
                           end_row=row - 1, end_column=1)

    for col, width in {'A': 26, 'B': 28, 'C': 36, 'D': 46}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:D{row - 1}'
    wb.save(OUT)

    people = row - 2
    merged = sum(1 for g in teams if len(g['members']) > 1)
    blanks = sum(1 for g in teams + solos for m in g['members']
                 if m['college'] == NOT_SPECIFIED)
    print(f'wrote {OUT}: {people} people ({len(teams)} teams, {len(solos)} solos), '
          f'{merged} merged team cells, {blanks} blank colleges')


if __name__ == '__main__':
    main()
